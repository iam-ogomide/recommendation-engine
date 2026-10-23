import argparse
import csv
import os
import sys
from datetime import date, datetime
from pathlib import Path
from typing import Any

from dotenv import load_dotenv
from openpyxl import load_workbook
from pymongo import MongoClient, ReplaceOne
from pymongo.errors import PyMongoError


PROJECT_DIR = Path(__file__).resolve().parent
INTEGER_FIELDS = {"cycle_day", "completed_cycles"}
DATE_FIELDS = {"subscribed_at", "last_payment_date", "next_run_date"}


def normalize_value(field: str, value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, str):
        value = value.strip()
        if not value:
            return None

    if field in INTEGER_FIELDS:
        try:
            return int(value)
        except (TypeError, ValueError) as exc:
            raise ValueError(f"Expected an integer in column '{field}', got {value!r}") from exc

    if field in DATE_FIELDS:
        if isinstance(value, datetime):
            return value
        if isinstance(value, date):
            return datetime.combine(value, datetime.min.time())
        if isinstance(value, str):
            try:
                return datetime.fromisoformat(value)
            except ValueError as exc:
                raise ValueError(f"Expected an ISO date in column '{field}', got {value!r}") from exc

    if field == "bvn":
        return str(value)
    return value


def records_from_csv(path: Path) -> list[dict[str, Any]]:
    with path.open("r", newline="", encoding="utf-8-sig") as file:
        reader = csv.DictReader(file)
        if not reader.fieldnames:
            raise ValueError(f"{path.name} is empty or has no header row")
        headers = [header.strip().lower() if header else "" for header in reader.fieldnames]
        if len(headers) != len(set(headers)) or any(not header for header in headers):
            raise ValueError("The input file must have unique, non-empty column headers")
        if "subscriber_id" not in headers:
            raise ValueError("The input file must contain a subscriber_id column")

        records = []
        for row_number, row in enumerate(reader, start=2):
            record = {
                header: normalize_value(header, row.get(original_header))
                for original_header, header in zip(reader.fieldnames, headers)
            }
            if any(value is not None for value in record.values()):
                if not record.get("subscriber_id"):
                    raise ValueError(f"Missing subscriber_id on row {row_number}")
                records.append(record)
        return records


def records_from_excel(path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)
        headers_row = next(rows, None)
        if not headers_row:
            raise ValueError(f"{path.name} is empty or has no header row")

        headers = [
            str(header).strip().lower() if header is not None else ""
            for header in headers_row
        ]
        if len(headers) != len(set(headers)) or any(not header for header in headers):
            raise ValueError("The input file must have unique, non-empty column headers")
        if "subscriber_id" not in headers:
            raise ValueError("The input file must contain a subscriber_id column")

        records = []
        for row_number, row in enumerate(rows, start=2):
            record = {
                header: normalize_value(header, row[index] if index < len(row) else None)
                for index, header in enumerate(headers)
            }
            if any(value is not None for value in record.values()):
                if not record.get("subscriber_id"):
                    raise ValueError(f"Missing subscriber_id on row {row_number}")
                records.append(record)
        return records
    finally:
        workbook.close()


def load_records(path: Path) -> list[dict[str, Any]]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return records_from_csv(path)
    if suffix in {".xlsx", ".xlsm"}:
        return records_from_excel(path)
    raise ValueError("Unsupported file type. Use a .csv, .xlsx, or .xlsm file.")


def main() -> int:
    load_dotenv(PROJECT_DIR / ".env")

    parser = argparse.ArgumentParser(
        description="Upload subscriber records from CSV or Excel to MongoDB."
    )
    parser.add_argument(
        "file",
        nargs="?",
        help="CSV or Excel file to upload (defaults to DATA_FILE in .env)",
    )
    args = parser.parse_args()

    uri = os.getenv("MONGODB_URI")
    database_name = os.getenv("MONGODB_DATABASE")
    collection_name = os.getenv("MONGODB_COLLECTION", "demo_subscribers_data")
    input_name = args.file or os.getenv("DATA_FILE", "subscribers_sim.csv")
    input_path = Path(input_name)
    if not input_path.is_absolute():
        input_path = PROJECT_DIR / input_path

    missing = [
        name
        for name, value in (
            ("MONGODB_URI", uri),
            ("MONGODB_DATABASE", database_name),
        )
        if not value or "<" in value or ">" in value
    ]
    if missing:
        print(
            "Set the following values in your .env file: " + ", ".join(missing),
            file=sys.stderr,
        )
        return 1
    if not input_path.is_file():
        print(f"Input file not found: {input_path}", file=sys.stderr)
        return 1

    try:
        records = load_records(input_path)
        if not records:
            print("No records found; nothing was uploaded.")
            return 0

        if any("subscriber_id" not in record for record in records):
            raise ValueError("The input file must contain a subscriber_id column")

        operations = [
            ReplaceOne({"subscriber_id": record["subscriber_id"]}, record, upsert=True)
            for record in records
        ]
        with MongoClient(uri, serverSelectionTimeoutMS=10000) as client:
            collection = client[database_name][collection_name]
            collection.create_index("subscriber_id", unique=True)
            result = collection.bulk_write(operations, ordered=True)

        print(
            f"Uploaded {len(records)} records to "
            f"{database_name}.{collection_name} "
            f"({result.upserted_count} inserted, {result.modified_count} updated)."
        )
        return 0
    except (OSError, ValueError, PyMongoError) as exc:
        print(f"Upload failed: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
